from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "QualiCoco_Plano_Editorial_Outubro_2026_Aprovacao.docx"

GREEN = "2F6B4F"
LIGHT_GREEN = "EAF2EC"
PALE = "F5F7F4"
GRAY = "D9DED9"
TEXT = "202522"
MUTED = "5F6B63"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=GRAY, sz='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_keep_with_next(paragraph, value=True):
    pPr = paragraph._p.get_or_add_pPr()
    keep = pPr.find(qn('w:keepNext'))
    if value and keep is None:
        pPr.append(OxmlElement('w:keepNext'))
    elif not value and keep is not None:
        pPr.remove(keep)


def set_cell_text(cell, text, bold=False, color=TEXT, size=9.2, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    set_cell_border(cell)


def add_run(p, text, bold=False, italic=False, color=TEXT, size=10.5):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def add_label_value(doc, label, value):
    p = doc.add_paragraph(style='Body Text')
    add_run(p, label + ' ', bold=True, color=GREEN, size=10.2)
    add_run(p, value, size=10.2)
    return p


def add_comment_line(doc, label='Comentários do cliente'):
    p = doc.add_paragraph(style='Comment Line')
    add_run(p, label + ': ', bold=True, color=MUTED, size=9.2)
    add_run(p, '______________________________________________________________', color='9AA39C', size=9.2)


def add_section_title(doc, text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    set_keep_with_next(p)
    return p


def add_idea(doc, item, video=False):
    p = doc.add_paragraph(item['title'], style='Heading 2')
    set_keep_with_next(p)
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Inches(1.45), Inches(5.55)]
    fields = [
        ('Data sugerida', item['date']),
        ('Tipo', 'Vídeo' if video else 'Post'),
        ('Tipo de produção', item['production']),
        ('Ideia', item['idea']),
        ('Defesa estratégica', item['defense']),
        ('Produto(s)', item['products']),
        ('Persona(s)', item['personas']),
        ('Rede(s)', item['networks']),
        ('Formato', item['format']),
        ('Campanha/Data', item['campaign']),
    ]
    for idx, (label, value) in enumerate(fields):
        cells = t.add_row().cells
        cells[0].width = widths[0]
        cells[1].width = widths[1]
        shade(cells[0], LIGHT_GREEN if idx % 2 == 0 else PALE)
        shade(cells[1], 'FFFFFF')
        set_cell_text(cells[0], label, bold=True, color=GREEN, size=8.8)
        set_cell_text(cells[1], value, size=9.1)
    doc.add_paragraph('', style='Small Space')
    add_comment_line(doc)
    doc.add_paragraph('', style='Small Space')


posts = [
    dict(title='01 — Café da manhã que resolve a semana', date='01/10', production='Post interno', idea='Carrossel com combinações de café da manhã e lanche usando diferentes possibilidades do portfólio de coco, mostrando como um mesmo território pode aparecer em momentos variados da rotina.', defense='Abre o mês com utilidade e repertório de consumo. Ajuda a marca a ocupar o território de praticidade sem transformar o conteúdo em catálogo.', products='Leite de coco, coco ralado, requeijão de coco ou outro item adequado às combinações.', personas='Letícia e Helena', networks='Instagram e Facebook', format='Carrossel no feed', campaign='Nenhuma; abertura editorial do mês.'),
    dict(title='02 — O coco além do óbvio', date='03/10', production='Post interno', idea='Carrossel de descoberta apresentando diferentes categorias do portfólio QualiCoco e seus contextos de uso, da cozinha ao consumo cotidiano.', defense='Amplia a percepção da marca e evidencia a arquitetura de portfólio. É uma oportunidade de diferenciação em relação à comunicação que trata coco apenas como ingrediente de sobremesa.', products='Coco ralado, leite de coco, óleo de coco, água de coco, açúcar de coco, farinha de coco e produtos de origem vegetal.', personas='Letícia, Helena e Júlia', networks='Instagram e Facebook', format='Carrossel no feed', campaign='Nenhuma.'),
    dict(title='03 — Lanche prático para a rotina corrida', date='06/10', production='Post interno', idea='Conteúdo de inspiração com uma montagem de lanche simples, transportável e adaptável à rotina de trabalho, estudo ou deslocamento.', defense='Conecta praticidade, ocasião e identificação. A marca aparece como parte de uma rotina possível, sem promessa funcional ou discurso clínico.', products='Água de coco integral em lata e requeijão de coco, conforme a composição aprovada.', personas='Helena e Júlia', networks='Instagram, Facebook e Stories', format='Imagem estática ou carrossel curto', campaign='Nenhuma.'),
    dict(title='04 — Receita de família com cara de casa', date='09/10', production='Post interno', idea='Conteúdo afetivo sobre uma preparação caseira em que o coco participa como ingrediente de memória, sabor e encontro.', defense='Equilibra a presença mais utilitária do início do mês com vínculo emocional. Reforça sabor e proximidade sem depender de uma data comemorativa.', products='Coco ralado e leite de coco.', personas='Letícia e Helena', networks='Instagram e Facebook', format='Imagem estática com sequência de detalhes ou carrossel', campaign='Nenhuma.'),
    dict(title='05 — Dia das Crianças: lanche que agrada todo mundo', date='12/10', production='Post interno', idea='Ideia de lanche ou sobremesa para compartilhar em família, com foco em preparo simples, visual convidativo e participação de diferentes gerações.', defense='Usa a data de forma conectada ao território de alimentação e convivência. A abordagem é de ocasião e afeto, sem transformar o conteúdo em comunicação infantilizada.', products='Coco ralado, leite de coco, calda de coco ou achocolatado, conforme a receita escolhida.', personas='Letícia e Helena', networks='Instagram, Facebook e Stories', format='Carrossel no feed', campaign='Dia das Crianças.'),
    dict(title='06 — O que faz sentido ler no rótulo?', date='16/10', production='Post interno', idea='Carrossel educativo sobre como observar informações de embalagem e escolher um produto adequado à ocasião de consumo, sem prescrição ou interpretação clínica.', defense='Marca a presença no Dia Mundial da Alimentação com educação e responsabilidade. Ajuda a QualiCoco a construir autoridade por clareza, não por claims não sustentados.', products='Categorias do portfólio que permitam comparação objetiva de formato, aplicação e ocasião.', personas='Júlia e Helena', networks='Instagram e Facebook', format='Carrossel educativo', campaign='Dia Mundial da Alimentação.'),
    dict(title='07 — Água de coco para acompanhar a rotina', date='18/10', production='Post interno', idea='Conteúdo de ocasião mostrando a água de coco integral em lata em situações de consumo fora de casa, como pós-atividade, deslocamento ou pausa do dia.', defense='Dá contexto a um produto específico e transforma sua presença em cena de uso. A comunicação deve permanecer em praticidade e rotina, sem promessas de saúde.', products='Água de coco integral em lata 269 mL.', personas='Helena e Júlia', networks='Instagram, Facebook e Stories', format='Imagem estática ou Reels de apoio no Stories', campaign='Nenhuma.'),
    dict(title='08 — Sobremesa simples com efeito de receita salvável', date='21/10', production='Post interno', idea='Carrossel de inspiração com uma sobremesa visualmente atraente, poucos momentos de preparo e indicação clara de ocasião de consumo.', defense='Prioriza salvamentos e compartilhamentos. O produto aparece integrado à receita, com benefício de uso demonstrado em vez de apresentação genérica.', products='Leite de coco, coco ralado e calda de coco.', personas='Letícia e Helena', networks='Instagram e Facebook', format='Carrossel passo a passo', campaign='Nenhuma.'),
    dict(title='09 — Rotina de trabalho pede lanche prático', date='28/10', production='Post interno', idea='Conteúdo de identificação sobre escolhas e organização de um lanche para o expediente, com foco em praticidade e pausa.', defense='Aproveita o Dia do Servidor Público sem limitar a mensagem ao funcionalismo. A data funciona como gancho para falar de trabalho, rotina e alimentação possível.', products='Água de coco integral em lata e produtos de origem vegetal adequados ao lanche.', personas='Helena e Júlia', networks='Instagram, Facebook e Stories', format='Carrossel ou imagem estática', campaign='Dia do Servidor Público, tratado como oportunidade de rotina.'),
    dict(title='10 — O sabor do coco em mais de uma ocasião', date='30/10', production='Post interno', idea='Fechamento do mês reunindo diferentes ocasiões de consumo e categorias do portfólio, da receita ao lanche e da cozinha ao consumo fora de casa.', defense='Consolida a amplitude da marca e fecha o calendário com uma visão de portfólio, sem repetir o formato de apresentação institucional genérica.', products='Seleção de categorias: coco ralado, leite de coco, óleo de coco, água de coco e produtos de origem vegetal.', personas='Letícia, Helena e Júlia', networks='Instagram e Facebook', format='Carrossel de fechamento', campaign='Nenhuma.'),
]

videos = [
    dict(title='11 — Parceiro(a) de receitas: panqueca ou bolo de lanche com coco', date='02/10', production='Vídeo — Parceiro(a) de Receitas', idea='Receita visualmente apetitosa para café da manhã ou lanche, com uso natural de coco ralado, leite de coco ou farinha de coco.', defense='Gera demonstração de uso, salvamentos e compartilhamentos. A presença do produto é orgânica porque resolve uma etapa real da receita.', products='Coco ralado, leite de coco ou farinha de coco, conforme a receita definida.', personas='Letícia e Helena', networks='Instagram e Facebook', format='Reels', campaign='Nenhuma.'),
    dict(title='12 — Parceiro(a) de receitas: sobremesa cremosa de fim de semana', date='10/10', production='Vídeo — Parceiro(a) de Receitas', idea='Receita de sobremesa com textura e montagem atraentes, pensada para ser reproduzida em casa e associada a um momento de compartilhamento.', defense='Explora sabor, afeto e apelo visual. Complementa o vídeo de lanche com uma ocasião de consumo distinta e aumenta a variedade do calendário.', products='Leite de coco, coco ralado e calda de coco, conforme a receita definida.', personas='Letícia e Helena', networks='Instagram e Facebook', format='Reels', campaign='Nenhuma.'),
    dict(title='13 — UGC Nutrição: lanche real de quem estuda nutrição', date='14/10', production='Vídeo — UGC Nutrição', idea='A estudante mostra, de forma espontânea, como organiza um lanche em um dia de estudos, comentando praticidade e preferência pessoal sem falar como profissional de saúde.', defense='Traz identificação e autenticidade para a marca. A escolha do perfil aproxima QualiCoco de uma rotina real, sem criar diagnóstico, prescrição ou promessa funcional.', products='Água de coco integral em lata e/ou produto de origem vegetal adequado ao lanche.', personas='Júlia e Helena', networks='Instagram e TikTok, se disponível', format='UGC vertical / Reels', campaign='Nenhuma.'),
    dict(title='14 — UGC Nutrição: o que vai na minha rotina corrida', date='23/10', production='Vídeo — UGC Nutrição', idea='Vídeo cotidiano acompanhando a estudante entre deslocamento, estudos e alimentação, com uma escolha QualiCoco inserida na rotina de maneira natural.', defense='Reforça a marca como companhia para ocasiões reais e amplia o repertório além da receita. O conteúdo deve preservar linguagem pessoal e evitar qualquer autoridade clínica.', products='Água de coco integral em lata, requeijão de coco ou outro produto aprovado para a ocasião.', personas='Júlia e Helena', networks='Instagram e TikTok, se disponível', format='UGC vertical / Reels', campaign='Nenhuma.'),
]

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, size, color in [('Title', 28, TEXT), ('Heading 1', 17, TEXT), ('Heading 2', 13, GREEN), ('Heading 3', 11, TEXT)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(6)

if 'Body Text' not in styles:
    styles.add_style('Body Text', WD_STYLE_TYPE.PARAGRAPH)
styles['Body Text'].font.name = 'Aptos'
styles['Body Text'].font.size = Pt(10.2)
styles['Body Text'].font.color.rgb = RGBColor.from_string(TEXT)
styles['Body Text'].paragraph_format.space_after = Pt(4)
styles['Body Text'].paragraph_format.line_spacing = 1.1

if 'Comment Line' not in styles:
    styles.add_style('Comment Line', WD_STYLE_TYPE.PARAGRAPH)
styles['Comment Line'].font.name = 'Aptos'
styles['Comment Line'].font.size = Pt(9.2)
styles['Comment Line'].paragraph_format.space_before = Pt(1)
styles['Comment Line'].paragraph_format.space_after = Pt(4)

if 'Small Space' not in styles:
    styles.add_style('Small Space', WD_STYLE_TYPE.PARAGRAPH)
styles['Small Space'].paragraph_format.space_after = Pt(2)
styles['Small Space'].paragraph_format.space_before = Pt(0)

# Header and footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_run(header, 'QUALICOCO  /  DOCUMENTO DE APROVAÇÃO', bold=True, color=GREEN, size=8.5)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(footer, 'Planejamento editorial de Social Media  •  Outubro de 2026', color=MUTED, size=8)

# Cover
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_run(p, 'Planejamento Editorial de Social Media', bold=True, color=TEXT, size=27)
p = doc.add_paragraph()
add_run(p, 'QualiCoco  |  Outubro de 2026', bold=True, color=GREEN, size=15)
p.paragraph_format.space_after = Pt(16)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'Documento de aprovação para o cliente', bold=True, color=TEXT, size=12)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'Este documento apresenta a proposta editorial do mês para validação estratégica. Os conteúdos ainda não incluem legendas finais, textos de arte, roteiros completos ou briefings de produção.', color=MUTED, size=10.5)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
for row, (label, value) in zip(meta.rows, [('Cliente', 'QualiCoco'), ('Marca', 'QualiCoco'), ('Período', 'Outubro de 2026'), ('Etapa', 'Aprovação do planejamento editorial')]):
    row.cells[0].width = Inches(1.55)
    row.cells[1].width = Inches(5.35)
    shade(row.cells[0], LIGHT_GREEN)
    shade(row.cells[1], 'FFFFFF')
    set_cell_text(row.cells[0], label, bold=True, color=GREEN, size=9.2)
    set_cell_text(row.cells[1], value, size=9.6)

doc.add_paragraph('', style='Small Space')
p = doc.add_paragraph(style='Body Text')
add_run(p, 'Como comentar', bold=True, color=GREEN, size=10.5)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'Use os comentários do Google Docs para aprovar, pedir ajustes ou registrar dúvidas. Ao final, há uma área de decisão geral para consolidar a aprovação do mês.', color=TEXT, size=10.2)
doc.add_page_break()

# Executive reading
add_section_title(doc, '1. Leitura Estratégica do Mês', 1)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'A proposta de outubro posiciona a QualiCoco como uma marca de portfólio amplo, capaz de participar de diferentes momentos de alimentação e de rotina. O mês combina descoberta de categorias, inspiração culinária, utilidade e identificação, com datas comemorativas usadas apenas quando ajudam a construir uma conversa verdadeira.', size=10.5)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'A seleção mantém a marca QualiCoco como único escopo. QualiCau não é considerada neste planejamento.', bold=True, color=GREEN, size=10.5)

add_section_title(doc, 'Direção editorial', 2)
for text in [
    'Mostrar o coco em mais de uma ocasião: café da manhã, lanche, receita, sobremesa, rotina fora de casa e alimentação cotidiana.',
    'Usar o produto dentro de um contexto de consumo, evitando apresentações genéricas de catálogo.',
    'Construir valor por sabor, praticidade, repertório e clareza, sem criar claims nutricionais ou funcionais não sustentados.',
    'Equilibrar conteúdos de marca, produto, inspiração, utilidade e relacionamento.',
]:
    p = doc.add_paragraph(style='Body Text')
    p.style = styles['List Bullet']
    add_run(p, text, size=10.2)

add_section_title(doc, 'Personas prioritárias', 2)
add_label_value(doc, 'Personas utilizadas:', 'Letícia, Helena e Júlia, conforme a base de conhecimento atual da QualiCoco.')
add_label_value(doc, 'Critério de uso:', 'As personas aparecem de acordo com comportamento, rotina e ocasião de consumo; não há obrigação de contemplar todas em cada conteúdo.')

add_section_title(doc, '2. Oportunidades de Outubro', 1)
opp = doc.add_table(rows=1, cols=3)
opp.alignment = WD_TABLE_ALIGNMENT.CENTER
opp.autofit = False
for i, h in enumerate(['Data', 'Oportunidade', 'Uso no planejamento']):
    opp.rows[0].cells[i].width = [Inches(1.0), Inches(2.0), Inches(4.0)][i]
    shade(opp.rows[0].cells[i], GREEN)
    set_cell_text(opp.rows[0].cells[i], h, bold=True, color='FFFFFF', size=9.0)
set_repeat_table_header(opp.rows[0])
for rowdata in [
    ('12/10', 'Dia das Crianças', 'Lanche e sobremesa para compartilhar em família, com foco em ocasião e afeto.'),
    ('15/10', 'Dia do Professor', 'Oportunidade secundária para conteúdos de rotina, estudo e pausa; não entrou como peça principal.'),
    ('16/10', 'Dia Mundial da Alimentação', 'Conteúdo educativo sobre leitura de rótulo e escolhas contextualizadas.'),
    ('28/10', 'Dia do Servidor Público', 'Gancho para falar de rotina de trabalho e lanche prático, sem restringir a comunicação ao funcionalismo.'),
]:
    cells = opp.add_row().cells
    for i, val in enumerate(rowdata):
        cells[i].width = [Inches(1.0), Inches(2.0), Inches(4.0)][i]
        shade(cells[i], PALE if len(opp.rows) % 2 == 0 else 'FFFFFF')
        set_cell_text(cells[i], val, size=9.0)

p = doc.add_paragraph(style='Body Text')
add_run(p, 'Nota de curadoria: ', bold=True, color=GREEN, size=9.8)
add_run(p, 'as datas funcionam como oportunidades, não como eixo obrigatório do calendário. O planejamento também considera sazonalidade de rotina, alimentação, estudos, trabalho e consumo fora de casa.', size=9.8)

doc.add_page_break()
add_section_title(doc, '3. Ideias de Posts', 1)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'As 10 ideias abaixo compõem o banco inicial de posts. A curadoria final aparece ao fim do documento.', color=MUTED, size=10.0)
for item in posts:
    add_idea(doc, item, video=False)

doc.add_page_break()
add_section_title(doc, '4. Ideias de Vídeos', 1)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'A proposta contempla exclusivamente os dois formatos solicitados: parceiro(a) de receitas e UGC de estudante de nutrição. Não são roteiros completos.', color=MUTED, size=10.0)
for item in videos:
    add_idea(doc, item, video=True)

add_section_title(doc, '5. Análise de Variedade do Conjunto', 1)
analysis = [
    ('Produtos e categorias', 'A distribuição passa por coco ralado, leite de coco, óleo de coco, água de coco, ingredientes culinários e produtos de origem vegetal. A recomendação é confirmar a SKU disponível antes da produção.'),
    ('Personas', 'Letícia, Helena e Júlia aparecem de forma distribuída, com maior presença de Helena e Júlia nas situações de rotina e de Letícia nas receitas e ocasiões de compartilhamento.'),
    ('Marca e produto', 'Há equilíbrio entre conteúdo de repertório de marca, aplicação de produto e cenas de consumo. O produto entra como solução dentro de uma situação, não apenas como objeto de apresentação.'),
    ('Formatos', 'O conjunto alterna carrossel, imagem estática, Reels, UGC vertical e Stories de apoio.'),
    ('Objetivos', 'Estão contemplados descoberta, utilidade, inspiração, salvamento, compartilhamento, identificação, relacionamento e demonstração de uso.'),
    ('Datas', 'As datas são pontuais e subordinadas à estratégia. O mês não depende de calendário comemorativo para fazer sentido.'),
    ('Diferenciação', 'A oportunidade central é explorar a amplitude do portfólio e as ocasiões de consumo, em vez de repetir uma comunicação restrita a receitas ou atributos genéricos de saudabilidade.'),
]
for label, text in analysis:
    p = doc.add_paragraph(style='Body Text')
    add_run(p, label + ': ', bold=True, color=GREEN, size=10.2)
    add_run(p, text, size=10.2)

add_section_title(doc, '6. Curadoria Recomendada 6 Posts + 2 Vídeos', 1)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'A seleção abaixo é a recomendação para o calendário principal. Ela combina abertura de mês, utilidade, data relevante, educação, inspiração, amplitude de portfólio e presença audiovisual.', size=10.3)
cur = doc.add_table(rows=1, cols=4)
cur.alignment = WD_TABLE_ALIGNMENT.CENTER
cur.autofit = False
for i, h in enumerate(['Data', 'Conteúdo', 'Formato', 'Por que entra']):
    cur.rows[0].cells[i].width = [Inches(0.75), Inches(2.25), Inches(1.0), Inches(3.0)][i]
    shade(cur.rows[0].cells[i], GREEN)
    set_cell_text(cur.rows[0].cells[i], h, bold=True, color='FFFFFF', size=8.8)
set_repeat_table_header(cur.rows[0])
for rowdata in [
    ('01/10', 'Café da manhã que resolve a semana', 'Post', 'Abre o mês com utilidade e repertório.'),
    ('06/10', 'Lanche prático para a rotina corrida', 'Post', 'Conecta ocasião, produto e identificação.'),
    ('12/10', 'Dia das Crianças: lanche que agrada todo mundo', 'Post', 'Usa uma data relevante com afeto e compartilhamento.'),
    ('16/10', 'O que faz sentido ler no rótulo?', 'Post', 'Constrói educação e responsabilidade.'),
    ('21/10', 'Sobremesa simples com efeito de receita salvável', 'Post', 'Tem alto potencial de salvamento e demonstração de uso.'),
    ('30/10', 'O sabor do coco em mais de uma ocasião', 'Post', 'Fecha o mês consolidando amplitude de portfólio.'),
    ('02/10', 'Parceiro(a): panqueca ou bolo de lanche com coco', 'Vídeo', 'Demonstração de uso e potencial de compartilhamento.'),
    ('14/10', 'UGC Nutrição: lanche real de quem estuda nutrição', 'Vídeo', 'Autenticidade e identificação com rotina real.'),
]:
    cells = cur.add_row().cells
    for i, val in enumerate(rowdata):
        cells[i].width = [Inches(0.75), Inches(2.25), Inches(1.0), Inches(3.0)][i]
        shade(cells[i], PALE if len(cur.rows) % 2 == 0 else 'FFFFFF')
        set_cell_text(cells[i], val, size=8.8)

add_section_title(doc, '7. Conteúdos Reserva', 1)
reserve = [
    ('03/10', 'O coco além do óbvio', 'Alternativa para reforçar descoberta de portfólio e diferenciação.'),
    ('09/10', 'Receita de família com cara de casa', 'Reserva afetiva caso o calendário precise de mais proximidade.'),
    ('10/10', 'Parceiro(a): sobremesa cremosa de fim de semana', 'Alternativa de vídeo com foco em sabor e ocasião.'),
    ('18/10', 'Água de coco para acompanhar a rotina', 'Oportunidade de destaque para produto específico e consumo fora de casa.'),
    ('23/10', 'UGC Nutrição: o que vai na minha rotina corrida', 'Reserva de UGC para ampliar identificação e rotina.'),
    ('28/10', 'Rotina de trabalho pede lanche prático', 'Oportunidade de data e contexto profissional.'),
]
res = doc.add_table(rows=1, cols=3)
res.alignment = WD_TABLE_ALIGNMENT.CENTER
res.autofit = False
for i, h in enumerate(['Data', 'Conteúdo', 'Função no calendário']):
    res.rows[0].cells[i].width = [Inches(0.8), Inches(2.8), Inches(3.4)][i]
    shade(res.rows[0].cells[i], GREEN)
    set_cell_text(res.rows[0].cells[i], h, bold=True, color='FFFFFF', size=9.0)
for rowdata in reserve:
    cells = res.add_row().cells
    for i, val in enumerate(rowdata):
        cells[i].width = [Inches(0.8), Inches(2.8), Inches(3.4)][i]
        shade(cells[i], PALE if len(res.rows) % 2 == 0 else 'FFFFFF')
        set_cell_text(cells[i], val, size=9.0)

add_section_title(doc, '8. Pontos de Atenção e Informações Ausentes', 1)
for text in [
    'Não foram identificadas campanhas sazonais, lançamentos ou ativações específicas de outubro de 2026 na base consultada.',
    'A base disponível contém o catálogo geral, mas não traz fichas individuais atualizadas para cada produto nem detalhes completos de embalagem. A seleção final de SKU deve ser confirmada antes da produção.',
    'Claims nutricionais, funcionais, certificações, ingredientes específicos e comparações de produto não devem ser adicionados sem respaldo documental.',
    'Os vídeos UGC devem apresentar a participante como estudante de nutrição, sem linguagem de nutricionista formada, diagnóstico, prescrição, recomendação clínica ou promessa de saúde.',
    'A comunicação permanece exclusivamente QualiCoco. QualiCau está fora do escopo deste documento.',
]:
    p = doc.add_paragraph(style='Body Text')
    p.style = styles['List Bullet']
    add_run(p, text, size=10.2)

doc.add_page_break()
add_section_title(doc, 'Decisão do cliente', 1)
p = doc.add_paragraph(style='Body Text')
add_run(p, 'Marque a decisão no Google Docs e registre ajustes nos comentários das ideias correspondentes.', size=10.2)
decision = doc.add_table(rows=3, cols=2)
decision.alignment = WD_TABLE_ALIGNMENT.LEFT
for row, (label, value) in zip(decision.rows, [
    ('Status', '☐ Aprovado  ☐ Aprovado com ajustes  ☐ Revisar proposta'),
    ('Ajustes gerais', '______________________________________________________________'),
    ('Responsável pelo retorno', '______________________________________________________________'),
]):
    row.cells[0].width = Inches(1.55)
    row.cells[1].width = Inches(5.35)
    shade(row.cells[0], LIGHT_GREEN)
    shade(row.cells[1], 'FFFFFF')
    set_cell_text(row.cells[0], label, bold=True, color=GREEN, size=9.2)
    set_cell_text(row.cells[1], value, size=9.2)

p = doc.add_paragraph(style='Body Text')
add_run(p, 'Base consultada: ', bold=True, color=GREEN, size=9.2)
add_run(p, 'client.md, brand.md, catalog.md, personas.md e competitors.md da base atual da QualiCoco. Não foram localizados documentos adicionais de campanhas, produtos ou embalagem além dos arquivos de estrutura.', color=MUTED, size=9.2)

doc.core_properties.title = 'Planejamento Editorial de Social Media QualiCoco Outubro 2026'
doc.core_properties.subject = 'Documento de aprovação para cliente'
doc.core_properties.author = 'Marketing OS'
doc.core_properties.keywords = 'QualiCoco, planejamento editorial, social media, outubro 2026'
doc.save(OUT)
print(OUT)
