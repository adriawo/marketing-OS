from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='PLANEJAMENTO EDITORIAL | SANTO ANTÔNIO | SETEMBRO 2026.docx'
NAVY='24384A'; LINE='D9D9D9'
def set_cell(c, text, bold=False, head=False, size=8):
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text); r.font.name='Rubik'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Rubik'); r.font.size=Pt(size); r.bold=bold
    if head:
        r.font.color.rgb=RGBColor(255,255,255); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),NAVY); c._tc.get_or_add_tcPr().append(sh); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(d, heads, rows, size=8):
    t=d.add_table(rows=1, cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,x in enumerate(heads): set_cell(t.rows[0].cells[i],x,True,True,size)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            set_cell(cells[i],x,False,False,size)
            if j%2:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F5F8F9'); cells[i]._tc.get_or_add_tcPr().append(sh)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
def h(d,x,level=1):
    p=d.add_paragraph(style='Heading 1' if level==1 else 'Heading 2'); p.paragraph_format.keep_with_next=True; r=p.add_run(x); r.font.name='Rubik'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Rubik'); return p
def para(d,x,italic=False):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.1; r=p.add_run(x); r.font.name='Rubik'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Rubik'); r.font.size=Pt(10.5); r.italic=italic

d=Document(); s=d.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7); s.top_margin=Cm(2.54); s.bottom_margin=Cm(2.54); s.left_margin=Cm(2.54); s.right_margin=Cm(2.54)
for name,size,bold in [('Normal',10.5,False),('Title',26,True),('Heading 1',24,True),('Heading 2',17,True)]:
    st=d.styles[name]; st.font.name='Rubik'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Rubik'); st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=hp.add_run('ENGENHO  |  PLANEJAMENTO EDITORIAL'); rr.font.name='Rubik'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('2E667A')
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=fp.add_run('Santo Antônio Alimentos  •  Em aprovação  •  Setembro 2026'); rr.font.name='Rubik'; rr.font.size=Pt(8)

p=d.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(110); p.add_run('PLANEJAMENTO EDITORIAL SMM')
p=d.add_paragraph(); r=p.add_run('Santo Antônio Alimentos | Setembro 2026'); r.font.name='Rubik'; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string('2E667A')
para(d,'Planejamento preparado em setembro para organizar oportunidades publicáveis em outubro de 2026. Propostas para discussão e aprovação antes da produção.')
table(d,['Marca','Período','Etapa'],[['Santo Antônio Alimentos','Setembro 2026','Em aprovação']],10)
d.add_page_break()
h(d,'Direcionamento do mês')
para(d,'Setembro é a etapa de planejamento; o calendário proposto antecipa outubro. A comunicação prioriza produto, desejo, ocasiões de consumo e reconhecimento das embalagens, equilibrando amendoim, doces tradicionais e evolução do portfólio. Não há campanha comercial aprovada na base; as datas abaixo são oportunidades editoriais, não ações promocionais.')
h(d,'Oportunidades de outubro')
table(d,['Data','Oportunidade','Relação com a marca'],[['Todo o mês','Outubro Rosa','Manifestação institucional interna de cuidado, sem associar produtos a alegações de saúde.'],['16 out.','Dia Mundial da Alimentação','Narrativa institucional sobre tradição e evolução do portfólio, sem claims nutricionais.'],['31 out.','Halloween','Oportunidade lúdica para snacks e doces, com foco em ocasião e sabor.']],9)
para(d,'Referências de data: Dia Mundial da Alimentação em 16 de outubro; Outubro Rosa como campanha anual de conscientização; Halloween em 31 de outubro.',True)

items=[
['01','01 out.','19h','Post','Post de produto','Qual sabor acompanha sua pausa? Composição de Amendoim Crocante e Amendoim Japonês.','Inicia o mês com produto e ocasião cotidiana, reforçando o amendoim como território da marca e o reconhecimento de embalagem.','Amendoim Crocante; Amendoim Japonês','Instagram + Facebook','Carrossel','—','Consumidor'],
['02','06 out.','19h','Post','Post de ocasião','Pasta de Amendoim Integral em sugestão visual de café da manhã, com foco em combinar e servir.','Mostra aplicação prática sem prometer benefício nutricional; aproxima a categoria da rotina.','Pasta de Amendoim Integral','Instagram + Facebook','Carrossel','—','Consumidor'],
['03','09 out.','12h','Post','Post de produto','Tradição que continua no presente: Paçoca e Pé de Moleque em composição editorial.','Ativa memória afetiva e conecta a história iniciada em 1983 à linguagem atual da marca.','Paçoca; Pé de Moleque','Instagram + Facebook','Imagem estática','Tradição','Consumidor'],
['04','13 out.','19h','Post','Post de produto','Pasta de Amendoim com Pistache: sabor e novidade do portfólio em linguagem contemporânea.','Dá visibilidade a lançamento identificado na base, sem criar claims.','Pasta de Amendoim com Pistache','Instagram + Facebook','Carrossel','Lançamento identificado','Consumidor'],
['05','20 out.','19h','Post','Post de produto','Lanche que cabe na pausa: Bananinha Balls com Chocolate Zero e Mariola Zero.','Trabalha praticidade e variedade da linha Zero, sem extrapolar características documentadas.','Bananinha Balls com Chocolate Zero; Mariola Zero','Instagram + Facebook','Carrossel','Linha Zero','Consumidor'],
['06','31 out.','18h','Post','Post de ocasião','Halloween de sabores: Amendoim Pralinê, Paçoca e Pé de Moleque para compartilhar.','Aproveita data cultural com aderência natural aos doces e snacks, mantendo produto central.','Amendoim Pralinê; Paçoca; Pé de Moleque','Instagram + Facebook','Carrossel','Halloween','Consumidor'],
['07','16 out.','10h','Post','Post institucional','No Dia Mundial da Alimentação: tradição, amendoim e evolução convivem na trajetória da empresa.','Fortalece posicionamento corporativo em data aderente ao setor, sem promessas técnicas.','Portfólio Santo Antônio Alimentos','LinkedIn','Imagem estática','Dia Mundial da Alimentação','LinkedIn'],
['08','27 out.','10h','Post','Post institucional','Comunicação em evolução: tornar produtos e embalagens mais reconhecíveis no ponto de venda.','Traduz para B2B a modernização da marca e o papel da comunicação na preferência.','Portfólio Santo Antônio Alimentos','LinkedIn','Carrossel PDF','Modernização da marca','LinkedIn'],
['09','15 out.','10h','Post','Material interno','Mensagem interna de Outubro Rosa com encaminhamento a fontes oficiais de saúde.','A campanha é pertinente ao público interno e evita associar alimentos a prevenção ou tratamento.','Não aplicável','Comunicação interna','E-flyer interno','Outubro Rosa','Materiais Internos'],
['10','29 out.','10h','Post','Material interno','O que a marca quer que seja lembrado? Tradição desde 1983, proximidade e produtos no cotidiano.','Reforça internamente a ligação entre comunicação, reconhecimento de marca e ponto de venda.','Portfólio Santo Antônio Alimentos','Comunicação interna','E-flyer interno','Reconhecimento de marca','Materiais Internos'],
['V1','03 out.','11h','Vídeo','Vídeo animado','Reel animado do pote à pausa: Pasta de Amendoim Integral no café da manhã e no lanche.','Formato dinâmico para tornar a aplicação do produto desejável e memorável.','Pasta de Amendoim Integral','Instagram + Facebook','Reel animado','—','Consumidor'],
['V2','31 out.','11h','Vídeo','Vídeo animado','Reel animado de Halloween: revelação lúdica de Amendoim Pralinê, Paçoca e Pé de Moleque.','A data favorece identificação e a animação mantém o foco nos sabores do portfólio.','Amendoim Pralinê; Paçoca; Pé de Moleque','Instagram + Facebook','Reel animado','Halloween','Consumidor']]
heads=['Nº','Data','Horário','Tipo','Produção','Ideia','Defesa','Produto(s)','Rede(s)','Formato','Campanha/Data']
for group in ['Consumidor','LinkedIn','Materiais Internos']:
    d.add_page_break(); h(d,group)
    for x in [a for a in items if a[-1]==group]:
        h(d,x[0]+' — '+x[5],2)
        table(d,['Campo','Informação'],[
            ['Data | Horário | Tipo',x[1]+' | '+x[2]+' | '+x[3]],
            ['Produção',x[4]],
            ['Ideia',x[5]],
            ['Defesa',x[6]],
            ['Produto(s)',x[7]],
            ['Rede(s) | Formato',x[8]+' | '+x[9]],
            ['Campanha/Data',x[10]],
            ['Status','Em aprovação']],9.0)

d.add_page_break(); h(d,'Curadoria recomendada')
para(d,'Calendário principal: 6 posts e 2 vídeos. A seleção combina produto, tradição, evolução, ocasião cultural e posicionamento institucional. Todos permanecem em aprovação.')
table(d,['Nº','Conteúdo','Data','Rede','Formato'],[['01','Qual sabor acompanha sua pausa?','01 out.','Instagram + Facebook','Carrossel'],['03','Tradição que continua no presente','09 out.','Instagram + Facebook','Imagem estática'],['04','Pasta de Amendoim com Pistache','13 out.','Instagram + Facebook','Carrossel'],['06','Halloween de sabores','31 out.','Instagram + Facebook','Carrossel'],['07','Tradição e evolução do portfólio','16 out.','LinkedIn','Imagem estática'],['08','Comunicação em evolução','27 out.','LinkedIn','Carrossel PDF'],['V1','Do pote à pausa','03 out.','Instagram + Facebook','Reel animado'],['V2','Halloween de sabores','31 out.','Instagram + Facebook','Reel animado']],9)
h(d,'Conteúdos reserva')
table(d,['Nº','Ideia','Produto(s)','Rede(s)','Motivo da reserva'],[['02','Pasta de Amendoim Integral no café da manhã','Pasta de Amendoim Integral','Instagram + Facebook','Alternativa de ocasião para substituição ou conteúdo adicional.'],['05','Lanche que cabe na pausa','Bananinha Balls com Chocolate Zero; Mariola Zero','Instagram + Facebook','Preserva diversidade da linha Zero para uso posterior.'],['09','Mensagem interna de Outubro Rosa','Não aplicável','Comunicação interna','Depende de alinhamento da frente interna.'],['10','O que a marca quer que seja lembrado?','Portfólio Santo Antônio Alimentos','Comunicação interna','Depende de alinhamento com equipes e representantes.']],8.5)
d.add_page_break()
h(d,'Validação do planejamento')
table(d,['Item','Conferência'],[['Quantidade','10 posts: 6 para Instagram/Facebook, 2 para LinkedIn e 2 materiais internos; além de 2 vídeos em reels animados.'],['Calendário principal','6 posts + 2 vídeos.'],['Reserva','4 posts, para uso posterior ou alinhamento interno.'],['Fonte de informações','Base exclusiva da Santo Antônio Alimentos: cliente, marca e catálogo. Nenhuma informação de outros clientes foi utilizada.'],['Status','Em aprovação — nenhuma proposta está liberada para produção.']],9)
para(d,'Próxima etapa: após aprovação, consolidar apenas os conteúdos aprovados no calendário editorial e então iniciar a produção.',True)
d.core_properties.title='PLANEJAMENTO EDITORIAL | SANTO ANTÔNIO | SETEMBRO 2026'; d.core_properties.author='Engenho Comunicação e Design'; d.save(OUT); print(OUT)
